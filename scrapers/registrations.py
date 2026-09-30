"""Top selling makes in Singapore from LTA's monthly new car registrations (table M03).

LTA publishes the table as a PDF with one row per make, importer type and fuel type, and one
half year per page group. The first number on a row is that half year's total, so summing the
rows by make gives the year to date. LTA does not publish registrations by model.
"""
from __future__ import annotations

import io
import json
import logging
import re
from datetime import date
from typing import Any, Optional

from scrapers.base import BaseScraper, RobotsDisallowed, ScraperUnavailable

log = logging.getLogger(__name__)

ROW = re.compile(r"^(?P<make>[A-Z][A-Z0-9 .&'/]*?) (?:AMD|PI) (?P<fuel>[A-Za-z][A-Za-z() \-]*?) (?P<total>\d+)(?: |$)")
# No word boundaries: the PDF runs month headers together, as in "2026-022026-01".
MONTH = re.compile(r"(20\d\d)-(0[1-9]|1[0-2])")


def parse_registrations(text: str) -> tuple[dict[str, dict[str, int]], list[str]]:
    """({make: {"total": n, "ev": n, "petrol": n}}, sorted months covered) from the table's text.

    petrol counts petrol and petrol hybrid rows, since most petrol cars sold now are hybrids.
    """
    makes: dict[str, dict[str, int]] = {}
    for line in text.splitlines():
        m = ROW.match(line.strip())
        if not m:
            continue
        make = m["make"].replace(".", "").strip()
        row = makes.setdefault(make, {"total": 0, "ev": 0, "petrol": 0})
        n = int(m["total"])
        row["total"] += n
        fuel = m["fuel"].strip()
        if fuel == "Electric":
            row["ev"] += n
        elif fuel.startswith("Petrol"):
            row["petrol"] += n
    months = sorted({f"{y}-{mo}" for y, mo in MONTH.findall(text)})
    return makes, months


class RegistrationsScraper(BaseScraper):
    name = "registrations"

    def pdf_text(self, url: str) -> str:
        # fetch() is text only, so the PDF is fetched here and its extracted text is cached instead.
        cached = self._read_cache(url)
        if cached is not None:
            return cached
        if not self._allowed(url):
            raise RobotsDisallowed(url)
        from pypdf import PdfReader
        from urllib.parse import urlsplit

        BaseScraper._throttle.wait(urlsplit(url).netloc)
        log.info("GET %s", url)
        resp = self.client.get(url)
        resp.raise_for_status()
        text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(resp.content)).pages)
        self._write_cache(url, text)
        return text

    def parse(self, text: str) -> list:
        makes, months = parse_registrations(text)
        if not makes:
            raise ScraperUnavailable("no make rows found in the LTA registrations table")
        return [makes, months]

    def run(self) -> list:
        return self.parse(self.pdf_text(self.cfg["sources"]["lta_registrations_by_make"]))


# Body type split, from the Excel copy of table M03 on LTA DataMall.
#
# In the PDF an empty cell simply disappears from the extracted text, so a count cannot be tied
# to its body type. In the spreadsheet every count keeps its column, and each column carries a
# body type code under its month heading: HB, SDN, MPV, STW, SUV and CPE/Conv, then that month's
# total. Summing every column that carries the same code gives the year to date per body type.

BODY_CODES = {"HB": "Hatchback", "SDN": "Sedan", "MPV": "MPV", "STW": "Wagon", "SUV": "SUV"}


def body_code(cell: Any) -> Optional[str]:
    """Body type for a header cell such as "HB", "SDN" or "CPE/ Conv", else None."""
    if not isinstance(cell, str):
        return None
    code = re.sub(r"[^A-Z/]", "", cell.upper())
    if code.startswith("CPE") or code.startswith("CONV"):
        return "Coupe"
    return BODY_CODES.get(code)


def normalise_make(make: str) -> str:
    return re.sub(r"\s+", " ", make.replace(".", "")).strip().upper()


def _as_int(cell: Any) -> int:
    if isinstance(cell, bool) or cell is None:
        return 0
    if isinstance(cell, (int, float)):
        return int(cell)
    digits = re.sub(r"[^\d]", "", str(cell))
    return int(digits) if digits else 0


def _months_in(rows: list[tuple]) -> set[str]:
    found: set[str] = set()
    for row in rows:
        for cell in row:
            if hasattr(cell, "year") and hasattr(cell, "month"):
                found.add(f"{cell.year:04d}-{cell.month:02d}")
            elif isinstance(cell, str):
                found.update(f"{y}-{m}" for y, m in MONTH.findall(cell))
    return found


def parse_body_types_rows(sheets: list[list[tuple]]) -> tuple[dict[str, dict[str, int]], list[str]]:
    """({body type: {make: electric cars registered}}, sorted months) from the sheets' rows.

    The header row is the first one holding at least three body type codes. Make and fuel come
    from the columns left of the first body type column; a make written once for a block of
    rows (a merged cell) carries down. Only rows whose fuel is exactly "Electric" count, so
    plug in hybrids stay out.
    """
    counts: dict[str, dict[str, int]] = {}
    months: set[str] = set()
    for rows in sheets:
        header_idx, columns = None, {}
        for i, row in enumerate(rows[:60]):
            mapping = {j: body_code(c) for j, c in enumerate(row) if body_code(c)}
            if len(set(mapping.values())) >= 3:
                header_idx, columns = i, mapping
                break
        if header_idx is None:
            continue
        months |= _months_in(rows[: header_idx + 1])
        first_body = min(columns)
        make_col = fuel_col = None
        for row in rows[: header_idx + 1]:
            for j, cell in enumerate(row[:first_body]):
                text = str(cell or "").lower()
                if "make" in text and make_col is None:
                    make_col = j
                if "fuel" in text and fuel_col is None:
                    fuel_col = j
        make_col = 0 if make_col is None else make_col
        current_make = None
        for row in rows[header_idx + 1:]:
            left = list(row[:first_body])
            if make_col < len(left) and isinstance(left[make_col], str) and left[make_col].strip():
                current_make = normalise_make(left[make_col])
            if not current_make or current_make.startswith("TOTAL"):
                continue
            if fuel_col is not None and fuel_col < len(left):
                fuel = str(left[fuel_col] or "").strip()
            else:
                fuel = next((str(c).strip() for c in reversed(left) if isinstance(c, str) and c.strip()), "")
            if fuel.lower() != "electric":
                continue
            for j, body in columns.items():
                n = _as_int(row[j]) if j < len(row) else 0
                if n:
                    bucket = counts.setdefault(body, {})
                    bucket[current_make] = bucket.get(current_make, 0) + n
    return counts, sorted(months)


def parse_body_types_xlsx(data: bytes) -> tuple[dict[str, dict[str, int]], list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = [list(ws.iter_rows(values_only=True)) for ws in wb.worksheets]
    finally:
        wb.close()
    return parse_body_types_rows(sheets)


class BodyTypeRegistrationsScraper(BaseScraper):
    name = "registrations_body"

    def parse(self, data: bytes) -> list:
        counts, months = parse_body_types_xlsx(data)
        if not counts:
            raise ScraperUnavailable("no electric rows with a body type found in the LTA spreadsheet")
        return [counts, months]

    def run(self) -> list:
        url = self.cfg["sources"]["lta_registrations_by_make_xlsx"]
        # The parsed result is cached as JSON for the day, since the page cache holds text only.
        cached = self._read_cache(url)
        if cached is not None:
            data = json.loads(cached)
            return [data["counts"], data["months"]]
        if not self._allowed(url):
            raise RobotsDisallowed(url)
        from urllib.parse import urlsplit

        BaseScraper._throttle.wait(urlsplit(url).netloc)
        log.info("GET %s", url)
        resp = self.client.get(url)
        resp.raise_for_status()
        counts, months = self.parse(resp.content)
        self._write_cache(url, json.dumps({"counts": counts, "months": months}))
        return [counts, months]


def scrape_body_types(cfg: dict[str, Any], user_agent: str, run_date: date, force: bool = False) -> Optional[tuple[dict[str, dict[str, int]], list[str]]]:
    """({body type: {make: EV registrations}}, months) or None when the spreadsheet is unavailable."""
    if not cfg["sources"].get("lta_registrations_by_make_xlsx"):
        return None
    scraper = BodyTypeRegistrationsScraper(cfg, user_agent, run_date, force)
    try:
        counts, months = scraper.run()
        return counts, months
    except Exception as exc:
        log.warning("registrations by body type unavailable: %s", exc)
        return None
    finally:
        scraper.close()


def scrape_registrations(cfg: dict[str, Any], user_agent: str, run_date: date, force: bool = False) -> Optional[tuple[dict[str, dict[str, int]], list[str]]]:
    scraper = RegistrationsScraper(cfg, user_agent, run_date, force)
    try:
        makes, months = scraper.run()
        return makes, months
    except Exception as exc:
        log.warning("registrations unavailable: %s", exc)
        return None
    finally:
        scraper.close()
